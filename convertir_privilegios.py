import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict
import re

def procesar_archivo_privilegios(ruta_archivo_txt, ruta_archivo_excel):
    """
    Convierte archivo TXT de privilegios en Excel organizado por grupos
    """
    
    # Estructura para almacenar datos por grupo
    grupos = defaultdict(lambda: {
        'codigo': '',
        'nombre': '',
        'usuarios': [],
        'formas': [],
        'reportes': [],
        'procesos': []
    })
    
    print("Leyendo archivo...")
    
    with open(ruta_archivo_txt, 'r', encoding='utf-8', errors='ignore') as f:
        for linea in f:
            linea = linea.strip()
            if not linea:
                continue
                
            partes = linea.split('|')
            
            # Verificar que sea una línea de GRUPO
            if partes[0] == 'GRUPO' and len(partes) > 3:
                codigo_grupo = partes[1]
                nombre_grupo = partes[2]
                
                # Inicializar grupo si no existe
                if not grupos[codigo_grupo]['codigo']:
                    grupos[codigo_grupo]['codigo'] = codigo_grupo
                    grupos[codigo_grupo]['nombre'] = nombre_grupo
                
                # Buscar índice de "OPCIONES POR GRUPO"
                idx_opciones = None
                for i, parte in enumerate(partes):
                    if parte == 'OPCIONES POR GRUPO':
                        idx_opciones = i
                        break
                
                if idx_opciones is not None:
                    # CASO 1: Línea de usuario con "OPCIONES POR GRUPO" al final
                    if idx_opciones >= 6:
                        cod_usuario = partes[6] if len(partes) > 6 else ''
                        nombre_usuario = partes[7] if len(partes) > 7 else ''
                        activo = partes[8] if len(partes) > 8 else ''
                        
                        # Validar que solo sea Y o N en Activo
                        if (cod_usuario and 
                            cod_usuario.strip() and 
                            cod_usuario != 'Cod Usuario' and
                            cod_usuario != 'Formas' and
                            activo.strip().upper() in ['Y', 'N']):
                            
                            grupos[codigo_grupo]['usuarios'].append({
                                'codigo': cod_usuario,
                                'nombre': nombre_usuario,
                                'activo': activo
                            })
                            print(f"  Usuario agregado al grupo {codigo_grupo}: {cod_usuario} - {nombre_usuario}")
                    
                    # CASO 2: Línea de opciones (Formas o Reportes)
                    if len(partes) > idx_opciones + 1:
                        siguiente = partes[idx_opciones + 1]
                        
                        # Si después de "OPCIONES POR GRUPO" viene "Formas"
                        if siguiente == 'Formas' and len(partes) > idx_opciones + 7:
                            datos_opcion = partes[idx_opciones + 8:]
                            datos_opcion = [d for d in datos_opcion if d.strip()]
                            
                            if len(datos_opcion) >= 2:
                                grupos[codigo_grupo]['formas'].append(datos_opcion)
                                print(f"  Forma agregada al grupo {codigo_grupo}: {datos_opcion[0]}")
                        
                        # Si después de "OPCIONES POR GRUPO" viene "Reportes"
                        elif siguiente == 'Reportes' and len(partes) > idx_opciones + 7:
                            datos_reporte = partes[idx_opciones + 8:]
                            datos_reporte = [d for d in datos_reporte if d.strip()]
                            
                            if len(datos_reporte) >= 2:
                                grupos[codigo_grupo]['reportes'].append(datos_reporte)
                                print(f"  Reporte agregado al grupo {codigo_grupo}: {datos_reporte[0]}")
                        
                        # Si después de "OPCIONES POR GRUPO" viene "Procesos"
                        elif siguiente == 'Procesos' and len(partes) > idx_opciones + 7:
                            datos_proceso = partes[idx_opciones + 8:]
                            datos_proceso = [d for d in datos_proceso if d.strip()]
                            
                            if len(datos_proceso) >= 2:
                                grupos[codigo_grupo]['procesos'].append(datos_proceso)
                                print(f"  Proceso agregado al grupo {codigo_grupo}: {datos_proceso[0]}")
                
                else:
                    # Sin "OPCIONES POR GRUPO" - línea de usuario normal
                    if len(partes) > 5:
                        cod_usuario = partes[3]
                        nombre_usuario = partes[4]
                        activo = partes[5]
                        
                        # Validar que solo sea Y o N en Activo
                        if (cod_usuario and 
                            cod_usuario.strip() and 
                            cod_usuario != 'Cod Usuario' and
                            activo.strip().upper() in ['Y', 'N']):
                            
                            grupos[codigo_grupo]['usuarios'].append({
                                'codigo': cod_usuario,
                                'nombre': nombre_usuario,
                                'activo': activo
                            })
                            print(f"  Usuario agregado al grupo {codigo_grupo}: {cod_usuario} - {nombre_usuario}")
    
    print(f"\nSe encontraron {len(grupos)} grupos")
    for codigo, datos in grupos.items():
        print(f"  Grupo {codigo}: {len(datos['usuarios'])} usuarios, {len(datos['formas'])} formas, {len(datos['reportes'])} reportes, {len(datos['procesos'])} procesos")
    
    # Crear archivo Excel
    print("\nCreando archivo Excel...")
    wb = Workbook()
    wb.remove(wb.active)
    
    # Estilos
    titulo_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    titulo_font = Font(bold=True, color="FFFFFF", size=12)
    seccion_fill = PatternFill(start_color="B4C7E7", end_color="B4C7E7", fill_type="solid")
    seccion_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True, size=10)
    
    for codigo_grupo, datos in sorted(grupos.items()):
        # Crear nombre de hoja
        nombre_hoja = f"{codigo_grupo} {datos['nombre']}"[:31]
        nombre_hoja = re.sub(r'[\\/*?:\[\]]', '', nombre_hoja)
        
        ws = wb.create_sheet(title=nombre_hoja)
        fila_actual = 1
        
        # TÍTULO DEL GRUPO
        cell = ws.cell(row=fila_actual, column=1, value=f"{datos['codigo']} {datos['nombre']}")
        cell.font = titulo_font
        cell.fill = titulo_fill
        ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=7)
        fila_actual += 1
        
        # SECCIÓN DE USUARIOS
        if datos['usuarios']:
            fila_actual += 1
            
            ws.cell(row=fila_actual, column=1, value="#").font = header_font
            ws.cell(row=fila_actual, column=1).fill = header_fill
            ws.cell(row=fila_actual, column=2, value="Cod Usuario").font = header_font
            ws.cell(row=fila_actual, column=2).fill = header_fill
            ws.cell(row=fila_actual, column=3, value="Nombre del Usuario").font = header_font
            ws.cell(row=fila_actual, column=3).fill = header_fill
            ws.cell(row=fila_actual, column=4, value="Activo").font = header_font
            ws.cell(row=fila_actual, column=4).fill = header_fill
            fila_actual += 1
            
            for idx, usuario in enumerate(datos['usuarios'], start=1):
                ws.cell(row=fila_actual, column=1, value=idx)
                ws.cell(row=fila_actual, column=2, value=usuario['codigo'])
                ws.cell(row=fila_actual, column=3, value=usuario['nombre'])
                ws.cell(row=fila_actual, column=4, value=usuario['activo'])
                fila_actual += 1
            
            fila_actual += 1
        
        # SECCIÓN DE FORMAS
        cell = ws.cell(row=fila_actual, column=1, value="OPCIONES POR GRUPO")
        cell.font = seccion_font
        cell.fill = seccion_fill
        ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=7)
        fila_actual += 1
        
        encabezados = ['Formas', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
        for col, encabezado in enumerate(encabezados, start=1):
            cell = ws.cell(row=fila_actual, column=col, value=encabezado)
            cell.font = header_font
            cell.fill = header_fill
        fila_actual += 1
        
        if datos['formas']:
            for forma_data in datos['formas']:
                for col, valor in enumerate(forma_data, start=1):
                    if col <= 7 and valor and valor.strip():
                        ws.cell(row=fila_actual, column=col, value=valor.strip())
                fila_actual += 1
        
        fila_actual += 1
        
        # SECCIÓN DE REPORTES
        encabezados_reportes = ['Reportes', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
        for col, encabezado in enumerate(encabezados_reportes, start=1):
            cell = ws.cell(row=fila_actual, column=col, value=encabezado)
            cell.font = header_font
            cell.fill = header_fill
        fila_actual += 1
        
        if datos['reportes']:
            for reporte_data in datos['reportes']:
                for col, valor in enumerate(reporte_data, start=1):
                    if col <= 7 and valor and valor.strip():
                        ws.cell(row=fila_actual, column=col, value=valor.strip())
                fila_actual += 1
        
        fila_actual += 1
        
        # SECCIÓN DE PROCESOS
        encabezados_procesos = ['Procesos', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
        for col, encabezado in enumerate(encabezados_procesos, start=1):
            cell = ws.cell(row=fila_actual, column=col, value=encabezado)
            cell.font = header_font
            cell.fill = header_fill
        fila_actual += 1
        
        if datos['procesos']:
            for proceso_data in datos['procesos']:
                for col, valor in enumerate(proceso_data, start=1):
                    if col <= 7 and valor and valor.strip():
                        ws.cell(row=fila_actual, column=col, value=valor.strip())
                fila_actual += 1
        
        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 18
    
    wb.save(ruta_archivo_excel)
    print(f"\n✓ Archivo Excel creado exitosamente: {ruta_archivo_excel}")
    print(f"✓ Total de hojas creadas: {len(grupos)}")

if __name__ == "__main__":
    archivo_entrada = "privilegios.txt"
    archivo_salida = "privilegios_organizados.xlsx"
    
    try:
        procesar_archivo_privilegios(archivo_entrada, archivo_salida)
    except FileNotFoundError:
        print(f"Error: No se encontró el archivo '{archivo_entrada}'")
        print("Asegúrate de que el archivo esté en la misma carpeta que este script")
    except Exception as e:
        print(f"Error al procesar el archivo: {str(e)}")
        import traceback
        traceback.print_exc()