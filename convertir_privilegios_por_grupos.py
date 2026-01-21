import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict
import re

# Mapeo de códigos de grupo a nombre de Excel
# IMPORTANTE: Un código puede pertenecer a múltiples coordinaciones (lista)
MAPEO_GRUPOS = {
    '001': ['SIN COORDINACION'],
    '002': ['COORDINACION TIC'],
    '003': ['GERENCIA'],
    '004': ['CONTABILIDAD'],
    '005': ['COORDINACION DE PLANEACION Y DESARROLLO'],
    '006': ['COORDINACION DE GESTION AL ASOCIADO'],
    '007': ['GERENCIA', 'COORDINACION TIC'],
    '008': ['SIN COORDINACION'],
    '009': ['COORDINACION DE GESTION HUMANA Y ADMINISTRATIVA'],
    '010': ['SIN COORDINACION'],
    '011': ['GERENCIA'],
    '012': ['COORDINACION DE GESTION HUMANA Y ADMINISTRATIVA'],
    '013': ['DIRECCION FINANCIERA'],
    '014': ['CONTABILIDAD'],
    '015': ['COORDINACION DE CARTERA'],
    '016': ['COORDINACION DE GESTION HUMANA Y ADMINISTRATIVA'],
    '017': ['SIN COORDINACION'],
    '018': ['COORDINACION DE GESTION AL ASOCIADO'],
    '019': ['COORDINACION DE BIENESTAR SOCIAL'],
    '020': ['DIRECCION FINANCIERA', 'CONTABILIDAD'],
    '021': ['DIRECCION FINANCIERA', 'COORDINACION DE OPERACIONES'],
    '022': ['GERENCIA'],
    '023': ['GERENCIA', 'COORDINACION DE GESTION HUMANA Y ADMINISTRATIVA'],
    '024': ['SIN COORDINACION'],
    '025': ['GERENCIA', 'COORDINACION DE PLANEACION Y DESARROLLO'],
    '026': ['GERENCIA', 'DIRECCION FINANCIERA'],
    '027': ['GERENCIA', 'COORDINACION DE BIENESTAR SOCIAL'],
    '028': ['SIN COORDINACION'],
    '029': ['SIN COORDINACION'],
    '030': ['SIN COORDINACION'],
    '031': ['COORDINACION DE GESTION AL ASOCIADO'],
    '032': ['SIN COORDINACION'],
    '033': ['SIN COORDINACION'],
    '034': ['SIN COORDINACION'],
    '035': ['SIN COORDINACION'],
    '036': ['SIN COORDINACION'],
    '037': ['SIN COORDINACION'],
    '038': ['SIN COORDINACION'],
    '039': ['COORDINACION DE GESTION AL ASOCIADO'],
    '040': ['DIRECCION FINANCIERA', 'COORDINACION DE CARTERA'],
    '041': ['COORDINACION DE OPERACIONES'],
    '042': ['COORDINACION DE GESTION AL ASOCIADO'],
    '043': ['SIN COORDINACION'],
    '044': ['GERENCIA', 'COORDINACION DE GESTION AL ASOCIADO'],
    '045': ['COORDINACION DE CARTERA'],
    '046': ['SIN COORDINACION'],
    '048': ['SIN COORDINACION'],
    '049': ['GERENCIA', 'COORDINACION DE RIESGOS'],
    '050': ['SIN COORDINACION'],
    '051': ['SIN COORDINACION'],
    '052': ['SIN COORDINACION'],
    '061': ['COORDINACION DE OPERACIONES'],
    '065': ['COORDINACION DE GESTION AL ASOCIADO'],
    '066': ['COORDINACION DE CARTERA'],
    '070': ['DIRECCION FINANCIERA'],
    '071': ['COORDINACION DE OPERACIONES'],
    '080': ['COORDINACION DE OPERACIONES'],
    '081': ['COORDINACION DE CARTERA'],
    '082': ['DIRECCION FINANCIERA'],
    '115': ['SIN COORDINACION'],
    'ADS': ['COORDINACION TIC'],
    'COM': ['SIN COORDINACION'],
    'CRE': ['SIN COORDINACION'],
    'INA': ['SIN COORDINACION'],
    'PAT': ['SIN COORDINACION'],
    'SOP': ['SIN COORDINACION']
}

def procesar_archivo_privilegios(ruta_archivo_txt, carpeta_salida='privilegios_excel'):
    """
    Convierte archivo TXT de privilegios en múltiples Excel organizados por coordinación
    """
    import os
    
    # Crear carpeta de salida si no existe
    if not os.path.exists(carpeta_salida):
        os.makedirs(carpeta_salida)
    
    # Estructura para almacenar datos por grupo
    grupos = defaultdict(lambda: {
        'codigo': '',
        'nombre': '',
        'usuarios': [],
        'formas': [],
        'reportes': [],
        'procesos': []
    })
    
    print("Leyendo archivo...")
    
    with open(ruta_archivo_txt, 'r', encoding='utf-8', errors='ignore') as f:
        for linea in f:
            linea = linea.strip()
            if not linea:
                continue
                
            partes = linea.split('|')
            
            if partes[0] == 'GRUPO' and len(partes) > 3:
                codigo_grupo = partes[1]
                nombre_grupo = partes[2]
                
                if not grupos[codigo_grupo]['codigo']:
                    grupos[codigo_grupo]['codigo'] = codigo_grupo
                    grupos[codigo_grupo]['nombre'] = nombre_grupo
                
                idx_opciones = None
                for i, parte in enumerate(partes):
                    if parte == 'OPCIONES POR GRUPO':
                        idx_opciones = i
                        break
                
                if idx_opciones is not None:
                    if idx_opciones >= 6:
                        cod_usuario = partes[6] if len(partes) > 6 else ''
                        nombre_usuario = partes[7] if len(partes) > 7 else ''
                        activo = partes[8] if len(partes) > 8 else ''
                        
                        if (cod_usuario and 
                            cod_usuario.strip() and 
                            cod_usuario != 'Cod Usuario' and
                            cod_usuario != 'Formas' and
                            activo.strip().upper() in ['Y', 'N']):
                            
                            grupos[codigo_grupo]['usuarios'].append({
                                'codigo': cod_usuario,
                                'nombre': nombre_usuario,
                                'activo': activo
                            })
                            print(f"  Usuario agregado al grupo {codigo_grupo}: {cod_usuario} - {nombre_usuario}")
                    
                    if len(partes) > idx_opciones + 1:
                        siguiente = partes[idx_opciones + 1]
                        
                        if siguiente == 'Formas' and len(partes) > idx_opciones + 7:
                            datos_opcion = partes[idx_opciones + 8:]
                            datos_opcion = [d for d in datos_opcion if d.strip()]
                            
                            if len(datos_opcion) >= 2:
                                grupos[codigo_grupo]['formas'].append(datos_opcion)
                                print(f"  Forma agregada al grupo {codigo_grupo}: {datos_opcion[0]}")
                        
                        elif siguiente == 'Reportes' and len(partes) > idx_opciones + 7:
                            datos_reporte = partes[idx_opciones + 8:]
                            datos_reporte = [d for d in datos_reporte if d.strip()]
                            
                            if len(datos_reporte) >= 2:
                                grupos[codigo_grupo]['reportes'].append(datos_reporte)
                                print(f"  Reporte agregado al grupo {codigo_grupo}: {datos_reporte[0]}")
                        
                        elif siguiente == 'Procesos' and len(partes) > idx_opciones + 7:
                            datos_proceso = partes[idx_opciones + 8:]
                            datos_proceso = [d for d in datos_proceso if d.strip()]
                            
                            if len(datos_proceso) >= 2:
                                grupos[codigo_grupo]['procesos'].append(datos_proceso)
                                print(f"  Proceso agregado al grupo {codigo_grupo}: {datos_proceso[0]}")
                
                else:
                    if len(partes) > 5:
                        cod_usuario = partes[3]
                        nombre_usuario = partes[4]
                        activo = partes[5]
                        
                        if (cod_usuario and 
                            cod_usuario.strip() and 
                            cod_usuario != 'Cod Usuario' and
                            activo.strip().upper() in ['Y', 'N']):
                            
                            grupos[codigo_grupo]['usuarios'].append({
                                'codigo': cod_usuario,
                                'nombre': nombre_usuario,
                                'activo': activo
                            })
                            print(f"  Usuario agregado al grupo {codigo_grupo}: {cod_usuario} - {nombre_usuario}")
    
    print(f"\nSe encontraron {len(grupos)} grupos")
    
    # Agrupar por coordinación
    grupos_por_coordinacion = defaultdict(dict)
    
    for codigo_grupo, datos in grupos.items():
        nombres_excel = MAPEO_GRUPOS.get(codigo_grupo, ['SIN COORDINACION'])
        
        # Un grupo puede aparecer en múltiples Excel
        for nombre_excel in nombres_excel:
            grupos_por_coordinacion[nombre_excel][codigo_grupo] = datos
            print(f"  Grupo {codigo_grupo} ({datos['nombre']}): {len(datos['usuarios'])} usuarios → Excel: {nombre_excel}")
    
    print(f"\nSe crearán {len(grupos_por_coordinacion)} archivos Excel")
    
    # Estilos
    titulo_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    titulo_font = Font(bold=True, color="FFFFFF", size=12)
    seccion_fill = PatternFill(start_color="B4C7E7", end_color="B4C7E7", fill_type="solid")
    seccion_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True, size=10)
    
    # Crear un Excel por coordinación
    for nombre_coordinacion, grupos_dict in grupos_por_coordinacion.items():
        print(f"\nCreando Excel para: {nombre_coordinacion}")
        
        wb = Workbook()
        wb.remove(wb.active)
        
        for codigo_grupo, datos in sorted(grupos_dict.items()):
            nombre_hoja = f"{codigo_grupo} {datos['nombre']}"[:31]
            nombre_hoja = re.sub(r'[\\/*?:\[\]]', '', nombre_hoja)
            
            ws = wb.create_sheet(title=nombre_hoja)
            fila_actual = 1
            
            # TÍTULO DEL GRUPO
            cell = ws.cell(row=fila_actual, column=1, value=f"{datos['codigo']} {datos['nombre']}")
            cell.font = titulo_font
            cell.fill = titulo_fill
            ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=7)
            fila_actual += 1
            
            # SECCIÓN DE USUARIOS
            if datos['usuarios']:
                fila_actual += 1
                
                ws.cell(row=fila_actual, column=1, value="#").font = header_font
                ws.cell(row=fila_actual, column=1).fill = header_fill
                ws.cell(row=fila_actual, column=2, value="Cod Usuario").font = header_font
                ws.cell(row=fila_actual, column=2).fill = header_fill
                ws.cell(row=fila_actual, column=3, value="Nombre del Usuario").font = header_font
                ws.cell(row=fila_actual, column=3).fill = header_fill
                ws.cell(row=fila_actual, column=4, value="Activo").font = header_font
                ws.cell(row=fila_actual, column=4).fill = header_fill
                fila_actual += 1
                
                for idx, usuario in enumerate(datos['usuarios'], start=1):
                    ws.cell(row=fila_actual, column=1, value=idx)
                    ws.cell(row=fila_actual, column=2, value=usuario['codigo'])
                    ws.cell(row=fila_actual, column=3, value=usuario['nombre'])
                    ws.cell(row=fila_actual, column=4, value=usuario['activo'])
                    fila_actual += 1
                
                fila_actual += 1
            
            # SECCIÓN DE FORMAS
            cell = ws.cell(row=fila_actual, column=1, value="OPCIONES POR GRUPO")
            cell.font = seccion_font
            cell.fill = seccion_fill
            ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=7)
            fila_actual += 1
            
            encabezados = ['Formas', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
            for col, encabezado in enumerate(encabezados, start=1):
                cell = ws.cell(row=fila_actual, column=col, value=encabezado)
                cell.font = header_font
                cell.fill = header_fill
            fila_actual += 1
            
            if datos['formas']:
                for forma_data in datos['formas']:
                    for col, valor in enumerate(forma_data, start=1):
                        if col <= 7 and valor and valor.strip():
                            ws.cell(row=fila_actual, column=col, value=valor.strip())
                    fila_actual += 1
            
            fila_actual += 1
            
            # SECCIÓN DE REPORTES
            encabezados_reportes = ['Reportes', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
            for col, encabezado in enumerate(encabezados_reportes, start=1):
                cell = ws.cell(row=fila_actual, column=col, value=encabezado)
                cell.font = header_font
                cell.fill = header_fill
            fila_actual += 1
            
            if datos['reportes']:
                for reporte_data in datos['reportes']:
                    for col, valor in enumerate(reporte_data, start=1):
                        if col <= 7 and valor and valor.strip():
                            ws.cell(row=fila_actual, column=col, value=valor.strip())
                    fila_actual += 1
            
            fila_actual += 1
            
            # SECCIÓN DE PROCESOS
            encabezados_procesos = ['Procesos', 'Insertar', 'Modificar', 'Eliminar', 'Consultar', 'Ejecuta Procesos', 'Otros Procesos']
            for col, encabezado in enumerate(encabezados_procesos, start=1):
                cell = ws.cell(row=fila_actual, column=col, value=encabezado)
                cell.font = header_font
                cell.fill = header_fill
            fila_actual += 1
            
            if datos['procesos']:
                for proceso_data in datos['procesos']:
                    for col, valor in enumerate(proceso_data, start=1):
                        if col <= 7 and valor and valor.strip():
                            ws.cell(row=fila_actual, column=col, value=valor.strip())
                    fila_actual += 1
            
            # Ajustar anchos de columna
            ws.column_dimensions['A'].width = 50
            ws.column_dimensions['B'].width = 18
            ws.column_dimensions['C'].width = 18
            ws.column_dimensions['D'].width = 18
            ws.column_dimensions['E'].width = 18
            ws.column_dimensions['F'].width = 18
            ws.column_dimensions['G'].width = 18
        
        # Guardar archivo
        nombre_archivo = f"{nombre_coordinacion}.xlsx"
        nombre_archivo = re.sub(r'[\\/*?:\[\]]', '', nombre_archivo)
        ruta_completa = f"{carpeta_salida}/{nombre_archivo}"
        wb.save(ruta_completa)
        print(f"  ✓ Creado: {ruta_completa} con {len(grupos_dict)} hojas")
    
    print(f"\n✓ Proceso completado exitosamente")
    print(f"✓ Total de archivos Excel creados: {len(grupos_por_coordinacion)}")
    print(f"✓ Ubicación: carpeta '{carpeta_salida}'")

if __name__ == "__main__":
    archivo_entrada = "privilegios.txt"
    
    try:
        procesar_archivo_privilegios(archivo_entrada)
    except FileNotFoundError:
        print(f"Error: No se encontró el archivo '{archivo_entrada}'")
        print("Asegúrate de que el archivo esté en la misma carpeta que este script")
    except Exception as e:
        print(f"Error al procesar el archivo: {str(e)}")
        import traceback
        traceback.print_exc()